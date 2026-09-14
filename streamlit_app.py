import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import duckdb
import pandas as pd
import streamlit as st

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# PAGE SETTINGS
# ============================================================

st.set_page_config(
    page_title="LYNX Cross Search",
    page_icon="🔎",
    layout="wide",
    initial_sidebar_state="collapsed"
)

DATABASE_FILE = "lynx_cross.duckdb"
LOGO_FILE = Path("assets/akita_lynx_logo.png")
DUBAI_TIMEZONE = ZoneInfo("Asia/Dubai")


# ============================================================
# CSS
# ============================================================

st.html(
    """
    <style>

    .stApp {
        background: #f3f3f3;
    }

    .block-container {
        max-width: 1500px;
        padding-top: 0rem;
        padding-bottom: 2rem;
    }

    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    header[data-testid="stHeader"] {
        background: transparent;
    }

    html, body {
        font-family: Arial, Helvetica, sans-serif;
    }


    /* HEADER */

    .lynx-header {
        background: linear-gradient(
            90deg,
            #3d3d3a 0%,
            #4d4d48 100%
        );

        min-height: 76px;

        display: flex;
        align-items: center;
        justify-content: space-between;

        padding: 0 28px;

        margin-left: -1rem;
        margin-right: -1rem;

        border-bottom: 3px solid #ef433d;
    }

    .lynx-header-left {
        display: flex;
        align-items: center;
        gap: 22px;
        min-width: 0;
    }

    .lynx-logo {
        height: 48px;
        width: auto;
        object-fit: contain;
        flex-shrink: 0;
    }

    .lynx-divider {
        width: 1px;
        height: 36px;
        background: #73736e;
        flex-shrink: 0;
    }

    .lynx-app-name {
        color: white;
        font-size: 16px;
        font-weight: 700;
        letter-spacing: 1.8px;
        white-space: nowrap;
    }

    .lynx-header-right {
        display: flex;
        align-items: center;
        gap: 14px;
        flex-shrink: 0;
    }

    .database-status {
        color: #e6e6e6;
        font-size: 12px;
        font-weight: 600;
        letter-spacing: .4px;
        white-space: nowrap;
    }

    .database-dot {
        display: inline-block;
        width: 8px;
        height: 8px;
        background: #69b96b;
        border-radius: 50%;
        margin-right: 7px;
    }

    .database-count {
        background: #ef433d;
        color: white;
        padding: 8px 14px;
        border-radius: 4px;
        font-size: 11px;
        font-weight: 700;
        letter-spacing: .5px;
        white-space: nowrap;
    }


    /* HERO */

    .lynx-hero {
        background: linear-gradient(
            110deg,
            #292927 0%,
            #3f3f3c 64%,
            #4a4a46 100%
        );

        padding: 32px 44px 34px 44px;

        margin-left: -1rem;
        margin-right: -1rem;

        position: relative;
        overflow: hidden;
    }

    .lynx-hero::after {
        content: "";
        position: absolute;
        width: 280px;
        height: 280px;
        right: -80px;
        top: -145px;
        border: 1px solid rgba(255,255,255,.06);
        border-radius: 50%;
    }

    .hero-small {
        color: #bcbcb9;
        font-size: 11px;
        font-weight: 700;
        letter-spacing: 2.2px;
        margin-bottom: 10px;
    }

    .hero-title {
        color: white;
        font-size: 34px;
        line-height: 1.08;
        font-weight: 700;
        margin-bottom: 10px;
    }

    .hero-title-red {
        color: #ef433d;
    }

    .hero-description {
        color: #d3d3d0;
        font-size: 15px;
        max-width: 760px;
        line-height: 1.45;
    }

    .hero-line {
        background: #ef433d;
        width: 150px;
        height: 3px;
        margin-top: 20px;
    }


    /* WORKSPACE */

    .workspace-title {
        margin-top: 24px;
        margin-bottom: 4px;
        color: #333330;
        font-size: 23px;
        font-weight: 700;
    }

    .workspace-subtitle {
        color: #777773;
        font-size: 13px;
        margin-bottom: 14px;
    }


    /* INPUTS */

    label {
        color: #454541 !important;
        font-weight: 600 !important;
    }

    .stTextArea textarea {
        background: white !important;
        color: #252525 !important;
        border: 1px solid #d0d0cc !important;
        border-radius: 4px !important;
        font-size: 14px !important;
        padding: 12px !important;
    }

    .stTextArea textarea:focus {
        border-color: #ef433d !important;
        box-shadow: 0 0 0 1px #ef433d !important;
    }

    div[data-baseweb="select"] > div {
        background: white !important;
        border-color: #d0d0cc !important;
        color: #333330 !important;
        border-radius: 4px !important;
    }

    .stTextInput input {
        background: white !important;
        color: #252525 !important;
        border: 1px solid #d0d0cc !important;
        border-radius: 4px !important;
    }

    .stTextInput input:focus {
        border-color: #ef433d !important;
        box-shadow: 0 0 0 1px #ef433d !important;
    }


    /* RADIO */

    div[role="radiogroup"] {
        gap: 10px;
    }

    div[role="radiogroup"] label {
        background: white;
        border: 1px solid #d1d1cd;
        border-radius: 4px;
        padding: 8px 16px;
        transition: .15s ease;
    }

    div[role="radiogroup"] label:hover {
        border-color: #ef433d;
    }


    /* BUTTONS */

    .stButton > button {
        background: linear-gradient(
            90deg,
            #ef433d,
            #f05a50
        ) !important;

        color: white !important;
        border: none !important;
        border-radius: 4px !important;

        min-height: 46px;

        font-size: 13px;
        font-weight: 700;
        letter-spacing: .8px;

        transition: .15s ease;
    }

    .stButton > button:hover {
        background: #dc3833 !important;
        transform: translateY(-1px);
        box-shadow: 0 5px 12px rgba(0,0,0,.12);
    }

    .stDownloadButton > button {
        background: #454541 !important;
        color: white !important;
        border: none !important;
        border-radius: 4px !important;
        min-height: 42px;
        font-weight: 700;
        font-size: 12px;
        letter-spacing: .6px;
    }

    .stDownloadButton > button:hover {
        background: #ef433d !important;
        color: white !important;
    }


    /* SEARCH TIPS */

    .search-tips {
        background: #454541;
        border-left: 4px solid #ef433d;
        color: #d7d7d4;
        padding: 14px 16px;
        border-radius: 3px;
        margin-top: 14px;
        font-size: 12px;
        line-height: 1.55;
    }

    .search-tips-title {
        color: white;
        font-size: 13px;
        font-weight: 700;
        margin-bottom: 6px;
    }


    /* METRICS */

    div[data-testid="stMetric"] {
        background: white;
        border: 1px solid #dededb;
        border-top: 3px solid #ef433d;
        border-radius: 4px;
        padding: 14px 18px;
        box-shadow: 0 2px 7px rgba(0,0,0,.04);
    }

    div[data-testid="stMetricLabel"] {
        color: #777773 !important;
        font-size: 12px;
    }

    div[data-testid="stMetricValue"] {
        color: #333330 !important;
        font-weight: 700;
    }


    /* DATAFRAME */

    div[data-testid="stDataFrame"] {
        background: white;
        border: 1px solid #d9d9d5;
        border-radius: 4px;
        overflow: hidden;
        box-shadow: 0 2px 7px rgba(0,0,0,.04);
    }


    /* RESULTS */

    .results-count {
        color: #777773;
        font-size: 12px;
        margin-top: 5px;
        margin-bottom: 8px;
    }


    /* FILE UPLOADER */

    section[data-testid="stFileUploaderDropzone"] {
        background: white !important;
        border: 1px dashed #bdbdb8 !important;
        border-radius: 4px !important;
    }


    /* UPLOAD INFO */

    .upload-ready {
        background: white;
        border-left: 4px solid #ef433d;
        border-top: 1px solid #ddd;
        border-right: 1px solid #ddd;
        border-bottom: 1px solid #ddd;
        padding: 11px 15px;
        border-radius: 3px;
        color: #454541;
        margin-top: 10px;
        margin-bottom: 12px;
        font-size: 13px;
    }


    /* SECTION TITLE */

    .section-title {
        display: flex;
        align-items: center;
        gap: 10px;
        margin-top: 24px;
        margin-bottom: 12px;
        color: #333330;
        font-size: 20px;
        font-weight: 700;
    }

    .section-red-line {
        width: 4px;
        height: 23px;
        background: #ef433d;
        border-radius: 2px;
    }


    /* FOOTER */

    .lynx-footer {
        margin-top: 32px;
        padding: 18px 0;
        border-top: 1px solid #d5d5d1;
        color: #8a8a86;
        text-align: center;
        font-size: 11px;
    }


    @media (max-width: 900px) {

        .lynx-header {
            padding: 12px 18px;
            flex-wrap: wrap;
            gap: 12px;
        }

        .lynx-header-right {
            gap: 8px;
        }

        .database-count {
            display: none;
        }

        .hero-title {
            font-size: 30px;
        }
    }

    </style>
    """
)


# ============================================================
# NORMALIZE NUMBER
# ============================================================

def normalize_number(number):

    if number is None:
        return ""

    text = str(number).strip()

    if not text or text.lower() == "nan":
        return ""

    return re.sub(
        r"[^A-Za-z0-9]",
        "",
        text
    ).upper()


# ============================================================
# SPLIT NUMBERS
# ============================================================

def split_numbers(text):

    if text is None:
        return []

    text = str(text).strip()

    if not text or text.lower() == "nan":
        return []

    parts = re.split(
        r"[,;/\n\r]+",
        text
    )

    result = []

    for part in parts:

        part = part.strip()

        if part:
            result.append(part)

    return result


# ============================================================
# FORMAT CUSTOMER NUMBERS
# ============================================================

def format_customer_numbers(numbers):

    return " / ".join(numbers)


# ============================================================
# PREPARE CUSTOMER ROWS
# ============================================================

def prepare_customer_rows(values):

    prepared = []

    for index, value in enumerate(values):

        if pd.isna(value):
            raw_value = ""
        else:
            raw_value = str(value).strip()

        prepared.append({
            "row_no": index + 1,
            "raw_value": raw_value
        })

    return prepared


# ============================================================
# BUILD SEARCH TABLE
# ============================================================

def build_search_table(prepared_rows):

    search_rows = []
    original_rows = []

    for item in prepared_rows:

        row_no = item["row_no"]
        raw_value = item["raw_value"]

        numbers = split_numbers(
            raw_value
        )

        customer_display = (
            format_customer_numbers(numbers)
            if numbers
            else raw_value
        )

        original_rows.append({
            "row_no": row_no,
            "customer_display": customer_display,
            "has_numbers": bool(numbers)
        })

        for original_number in numbers:

            search_number = normalize_number(
                original_number
            )

            if search_number:

                search_rows.append({
                    "row_no": row_no,
                    "customer_display": customer_display,
                    "search_number": search_number
                })

    search_df = pd.DataFrame(
        search_rows,
        columns=[
            "row_no",
            "customer_display",
            "search_number"
        ]
    )

    original_df = pd.DataFrame(
        original_rows,
        columns=[
            "row_no",
            "customer_display",
            "has_numbers"
        ]
    )

    return search_df, original_df


# ============================================================
# SEARCH DATABASE
# ============================================================

def search_crosses(
    prepared_rows,
    source_filter
):

    search_df, original_df = build_search_table(
        prepared_rows
    )

    connection = duckdb.connect(
        DATABASE_FILE,
        read_only=True
    )

    try:

        if search_df.empty:

            matches = pd.DataFrame(
                columns=[
                    "row_no",
                    "customer_display",
                    "lynx_number",
                    "description",
                    "source_type"
                ]
            )

        else:

            connection.register(
                "search_inputs",
                search_df
            )

            if source_filter == "All":

                query = """
                    SELECT
                        s.row_no,
                        s.customer_display,
                        c.lynx_number,
                        c.description,
                        c.source_type
                    FROM search_inputs s
                    INNER JOIN crosses c
                        ON c.search_number = s.search_number
                    ORDER BY
                        s.row_no,
                        c.lynx_number
                """

                matches = connection.execute(
                    query
                ).fetchdf()

            else:

                query = """
                    SELECT
                        s.row_no,
                        s.customer_display,
                        c.lynx_number,
                        c.description,
                        c.source_type
                    FROM search_inputs s
                    INNER JOIN crosses c
                        ON c.search_number = s.search_number
                    WHERE c.source_type = ?
                    ORDER BY
                        s.row_no,
                        c.lynx_number
                """

                matches = connection.execute(
                    query,
                    [source_filter]
                ).fetchdf()

    finally:

        connection.close()


    results = []


    for _, original_row in original_df.iterrows():

        row_no = int(
            original_row["row_no"]
        )

        customer_display = (
            original_row["customer_display"]
        )

        has_numbers = bool(
            original_row["has_numbers"]
        )


        if not has_numbers:

            results.append({
                "No.": row_no,
                "Customer Part Number": customer_display,
                "LYNXauto Number": "",
                "Description": "",
                "Cross Type": "",
                "Status": ""
            })

            continue


        if matches.empty:

            row_matches = pd.DataFrame()

        else:

            row_matches = matches[
                matches["row_no"] == row_no
            ].copy()


        if row_matches.empty:

            results.append({
                "No.": row_no,
                "Customer Part Number": customer_display,
                "LYNXauto Number": "",
                "Description": "",
                "Cross Type": "",
                "Status": "Not Found"
            })

            continue


        row_matches = row_matches.drop_duplicates(
            subset=[
                "lynx_number",
                "description",
                "source_type"
            ]
        )


        for _, match in row_matches.iterrows():

            results.append({
                "No.": row_no,

                "Customer Part Number":
                    customer_display,

                "LYNXauto Number":
                    str(match["lynx_number"]).strip(),

                "Description":
                    str(match["description"]).strip(),

                "Cross Type":
                    str(match["source_type"]).strip(),

                "Status":
                    "Found"
            })


    return pd.DataFrame(
        results
    )


# ============================================================
# ADD MANUAL CROSS
# ============================================================
def get_lynx_description(lynx_number):

    lynx_number = str(
        lynx_number
    ).strip().upper()

    if not lynx_number:
        return ""

    connection = duckdb.connect(
        DATABASE_FILE,
        read_only=True
    )

    try:

        result = connection.execute(
            """
            SELECT description
            FROM crosses
            WHERE UPPER(TRIM(lynx_number)) = ?
              AND description IS NOT NULL
              AND TRIM(description) <> ''
            LIMIT 1
            """,
            [lynx_number]
        ).fetchone()

    finally:

        connection.close()

    if result:
        return str(result[0]).strip()

    return ""


def update_manual_description():

    lynx_number = st.session_state.get(
        "manual_lynx_number",
        ""
    )

    description = get_lynx_description(
        lynx_number
    )

    if description:

        st.session_state[
            "manual_description"
        ] = description

        st.session_state[
            "manual_description_found"
        ] = True

    else:

        st.session_state[
            "manual_description"
        ] = ""

        st.session_state[
            "manual_description_found"
        ] = False

def add_manual_cross(
    df,
    row_no,
    lynx_number,
    description=""
):

    updated_df = df.copy()

    lynx_number = str(
        lynx_number
    ).strip().upper()

    description = str(
        description
    ).strip()


    if not lynx_number:

        return (
            updated_df,
            False,
            "Please enter LYNXauto Number."
        )


    selected_rows = updated_df[
        updated_df["No."] == row_no
    ]


    if selected_rows.empty:

        return (
            updated_df,
            False,
            "Customer position was not found."
        )


    customer_part_number = str(
        selected_rows.iloc[0][
            "Customer Part Number"
        ]
    )


    existing_numbers = (
        selected_rows[
            "LYNXauto Number"
        ]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )


    if lynx_number in existing_numbers.values:

        return (
            updated_df,
            False,
            f"{lynx_number} is already added "
            "to this customer position."
        )


    # Remove Not Found placeholder
    updated_df = updated_df[
        ~(
            (updated_df["No."] == row_no)
            &
            (updated_df["Status"] == "Not Found")
        )
    ].copy()


    manual_row = pd.DataFrame(
        [
            {
                "No.": row_no,
                "Customer Part Number":
                    customer_part_number,
                "LYNXauto Number":
                    lynx_number,
                "Description":
                    description,
                "Cross Type":
                    "Manual",
                "Status":
                    "Found"
            }
        ]
    )


    updated_df = pd.concat(
        [
            updated_df,
            manual_row
        ],
        ignore_index=True
    )


    updated_df = (
        updated_df
        .sort_values(
            by=["No."],
            kind="stable"
        )
        .reset_index(drop=True)
    )


    return (
        updated_df,
        True,
        f"{lynx_number} added successfully."
    )


# ============================================================
# REMOVE MANUAL CROSS
# ============================================================

def remove_manual_cross(
    df,
    row_no,
    lynx_number
):

    updated_df = df.copy()

    lynx_number = str(
        lynx_number
    ).strip().upper()


    target_rows = updated_df[
        (updated_df["No."] == row_no)
        &
        (
            updated_df["LYNXauto Number"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            == lynx_number
        )
        &
        (
            updated_df["Cross Type"]
            .fillna("")
            .astype(str)
            .str.strip()
            == "Manual"
        )
    ]


    if target_rows.empty:

        return (
            updated_df,
            False,
            "Manual cross was not found."
        )


    customer_part_number = str(
        target_rows.iloc[0][
            "Customer Part Number"
        ]
    )


    # Remove selected manual cross
    updated_df = updated_df[
        ~(
            (updated_df["No."] == row_no)
            &
            (
                updated_df["LYNXauto Number"]
                .fillna("")
                .astype(str)
                .str.strip()
                .str.upper()
                == lynx_number
            )
            &
            (
                updated_df["Cross Type"]
                .fillna("")
                .astype(str)
                .str.strip()
                == "Manual"
            )
        )
    ].copy()


    # Check if anything remains for this customer position
    remaining_rows = updated_df[
        updated_df["No."] == row_no
    ]


    # If nothing remains, restore Not Found
    if remaining_rows.empty:

        not_found_row = pd.DataFrame(
            [
                {
                    "No.": row_no,
                    "Customer Part Number":
                        customer_part_number,
                    "LYNXauto Number":
                        "",
                    "Description":
                        "",
                    "Cross Type":
                        "",
                    "Status":
                        "Not Found"
                }
            ]
        )


        updated_df = pd.concat(
            [
                updated_df,
                not_found_row
            ],
            ignore_index=True
        )


    updated_df = (
        updated_df
        .sort_values(
            by=["No."],
            kind="stable"
        )
        .reset_index(drop=True)
    )


    return (
        updated_df,
        True,
        f"{lynx_number} removed successfully."
    )


# ============================================================
# EXCEL EXPORT
# ============================================================

def create_excel_download(
    df,
    sales_manager,
    reference_no,
    report_datetime
):

    output = BytesIO()

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Cross References"

    # Hide Excel gridlines
    worksheet.sheet_view.showGridLines = False


    # ========================================================
    # COLORS
    # ========================================================

    dark_fill = PatternFill(
        "solid",
        fgColor="202427"
    )

    red_fill = PatternFill(
        "solid",
        fgColor="E30613"
    )

    label_fill = PatternFill(
        "solid",
        fgColor="F2F2F2"
    )

    found_fill = PatternFill(
        "solid",
        fgColor="E2F0D9"
    )

    not_found_fill = PatternFill(
        "solid",
        fgColor="FCE4D6"
    )

    label_font = Font(
        color="222222",
        bold=True,
        size=11
    )

    thin_border = Border(
        left=Side(
            style="thin",
            color="D0D0D0"
        ),
        right=Side(
            style="thin",
            color="D0D0D0"
        ),
        top=Side(
            style="thin",
            color="D0D0D0"
        ),
        bottom=Side(
            style="thin",
            color="D0D0D0"
        )
    )


    # ========================================================
    # COLUMN WIDTHS
    # ========================================================

    worksheet.column_dimensions["A"].width = 11
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 22
    worksheet.column_dimensions["D"].width = 36
    worksheet.column_dimensions["E"].width = 21
    worksheet.column_dimensions["F"].width = 18


    # ========================================================
    # ROW HEIGHTS
    # ========================================================

    worksheet.row_dimensions[1].height = 48
    worksheet.row_dimensions[2].height = 44
    worksheet.row_dimensions[3].height = 38
    worksheet.row_dimensions[4].height = 8
    worksheet.row_dimensions[5].height = 8
    worksheet.row_dimensions[6].height = 26
    worksheet.row_dimensions[7].height = 26
    worksheet.row_dimensions[8].height = 8
    worksheet.row_dimensions[9].height = 30


    # ========================================================
    # LOGO
    # ========================================================

    if LOGO_FILE.exists():

        logo = XLImage(
            str(LOGO_FILE)
        )

        logo.width = 320
        logo.height = 112

        worksheet.add_image(
            logo,
            "A1"
        )


    # ========================================================
    # TITLE
    # ========================================================

    worksheet.merge_cells(
        "D1:F2"
    )

    worksheet["D1"] = (
        "CROSS REFERENCES"
    )

    worksheet["D1"].font = Font(
        bold=True,
        size=24,
        color="202427"
    )

    worksheet["D1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    worksheet.merge_cells(
        "D3:F3"
    )

    worksheet["D3"] = (
        "AKITA Spare Parts Trading • "
        "LYNX Spare Parts"
    )

    worksheet["D3"].font = Font(
        size=11,
        color="555555"
    )

    worksheet["D3"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    # ========================================================
    # RED DIVIDER
    # ========================================================

    for column in range(1, 7):

        worksheet.cell(
            row=4,
            column=column
        ).fill = red_fill


    # ========================================================
    # REPORT INFO
    # ========================================================

    worksheet["A6"] = "Sales Manager"

    worksheet.merge_cells(
        "B6:C6"
    )

    worksheet["B6"] = sales_manager


    worksheet["D6"] = "Date"

    worksheet.merge_cells(
        "E6:F6"
    )

    worksheet["E6"] = (
        report_datetime.strftime(
            "%d %B %Y"
        )
    )


    worksheet["A7"] = "Reference"

    worksheet.merge_cells(
        "B7:C7"
    )

    worksheet["B7"] = reference_no


    worksheet["D7"] = "Prepared by"

    worksheet.merge_cells(
        "E7:F7"
    )

    worksheet["E7"] = (
        "Technical Product Department"
    )


    info_ranges = [
        "A6:C6",
        "D6:F6",
        "A7:C7",
        "D7:F7"
    ]


    for range_ref in info_ranges:

        for row in worksheet[
            range_ref
        ]:

            for cell in row:

                cell.border = thin_border

                cell.alignment = Alignment(
                    vertical="center"
                )


    for cell_ref in [
        "A6",
        "D6",
        "A7",
        "D7"
    ]:

        cell = worksheet[
            cell_ref
        ]

        cell.fill = label_fill
        cell.font = label_font


    # ========================================================
    # TABLE HEADER
    # ========================================================

    headers = [
        "No.",
        "Customer Part Number",
        "LYNXauto Number",
        "Description",
        "Cross Type",
        "Status"
    ]


    for column_index, header in enumerate(
        headers,
        start=1
    ):

        cell = worksheet.cell(
            row=9,
            column=column_index,
            value=header
        )

        cell.fill = dark_fill

        cell.font = Font(
            color="FFFFFF",
            bold=True,
            size=11
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border


    # ========================================================
    # DATA
    # ========================================================

    start_row = 10


    for row_index, (_, row) in enumerate(
        df.iterrows(),
        start=start_row
    ):

        values = [
            row["No."],
            row["Customer Part Number"],
            row["LYNXauto Number"],
            row["Description"],
            row["Cross Type"],
            row["Status"]
        ]


        for column_index, value in enumerate(
            values,
            start=1
        ):

            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )


        # LYNX hyperlink

        lynx_number = str(
            row["LYNXauto Number"]
        ).strip()


        if lynx_number:

            lynx_cell = worksheet.cell(
                row=row_index,
                column=3
            )

            lynx_cell.hyperlink = (
                "https://ecatalogue."
                "lynxauto.jp/"
                f"{lynx_number}.html"
            )

            lynx_cell.font = Font(
                color="0563C1",
                bold=True,
                underline="single"
            )

            lynx_cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


        # Status

        status = str(
            row["Status"]
        ).strip()


        status_cell = worksheet.cell(
            row=row_index,
            column=6
        )


        if status == "Found":

            status_cell.fill = (
                found_fill
            )

            status_cell.font = Font(
                color="375623",
                bold=True
            )


        elif status == "Not Found":

            status_cell.fill = (
                not_found_fill
            )

            status_cell.font = Font(
                color="C00000",
                bold=True
            )


        status_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


        worksheet.cell(
            row=row_index,
            column=1
        ).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


        worksheet.cell(
            row=row_index,
            column=5
        ).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # ========================================================
    # EXCEL TABLE
    # ========================================================

    last_row = (
        start_row
        + len(df)
        - 1
    )


    if last_row >= start_row:

        table = Table(
            displayName="CrossReferencesTable",
            ref=f"A9:F{last_row}"
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style

        worksheet.add_table(
            table
        )


    # ========================================================
    # FREEZE
    # ========================================================

    worksheet.freeze_panes = "A10"


    # ========================================================
    # PRINT SETTINGS
    # ========================================================

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0


    if last_row >= start_row:

        worksheet.print_area = (
            f"A1:F{last_row}"
        )

    else:

        worksheet.print_area = "A1:F10"


    # ========================================================
    # SAVE
    # ========================================================

    workbook.save(
        output
    )

    output.seek(0)

    return output.getvalue()


# ============================================================
# RESULT FILTER
# ============================================================

def filter_results(
    df,
    status_filter,
    search_text
):

    filtered = df.copy()


    if status_filter != "All":

        filtered = filtered[
            filtered["Status"] == status_filter
        ]


    search_text = str(
        search_text
    ).strip().lower()


    if search_text:

        searchable_columns = [
            "Customer Part Number",
            "LYNXauto Number",
            "Description",
            "Cross Type"
        ]

        mask = pd.Series(
            False,
            index=filtered.index
        )


        for column in searchable_columns:

            mask = mask | (
                filtered[column]
                .fillna("")
                .astype(str)
                .str.lower()
                .str.contains(
                    search_text,
                    regex=False
                )
            )


        filtered = filtered[
            mask
        ]


    return filtered


# ============================================================
# BUILD DISPLAY TABLE
# ============================================================

def build_display_table(df):

    display_df = df.copy()


    display_df["LYNXauto Number"] = (
        display_df["LYNXauto Number"]
        .fillna("")
        .astype(str)
        .apply(
            lambda value: (
                f"https://ecatalogue.lynxauto.jp/{value.strip()}.html"
                if value.strip()
                else ""
            )
        )
    )


    return display_df


# ============================================================
# DATABASE CHECK
# ============================================================

database_path = Path(
    DATABASE_FILE
)


if not database_path.exists():

    st.error(
        "Database file lynx_cross.duckdb was not found."
    )

    st.stop()


# ============================================================
# HEADER
# ============================================================

st.html(
    """
    <div class="lynx-header">

        <div class="lynx-header-left">

            <img
                class="lynx-logo"
                src="https://lynxauto.jp/assets/img/logo.png"
            >

            <div class="lynx-divider"></div>

            <div class="lynx-app-name">
                CROSS SEARCH
            </div>

        </div>

        <div class="lynx-header-right">

            <div class="database-status">
                <span class="database-dot"></span>
                DATABASE ONLINE
            </div>

            <div class="database-count">
                1.37M REFERENCES
            </div>

        </div>

    </div>

    <div class="lynx-hero">

        <div class="hero-small">
            LYNXAUTO PRODUCT DATABASE
        </div>

        <div class="hero-title">
            Find Your
            <span class="hero-title-red">
                Cross Reference
            </span>
        </div>

        <div class="hero-description">
            Fast search across OEM and Aftermarket cross references
            in the LYNXauto product database.
        </div>

        <div class="hero-line"></div>

    </div>
    """
)


# ============================================================
# WORKSPACE
# ============================================================

st.html(
    """
    <div class="workspace-title">
        Cross Reference Search
    </div>

    <div class="workspace-subtitle">
        Enter customer part numbers manually or upload an Excel customer request.
    </div>
    """
)


# ============================================================
# SALES MANAGER
# ============================================================

sales_manager = st.text_input(
    "Sales Manager",
    placeholder="Enter Sales Manager name...",
    help=(
        "This name will be shown "
        "in the exported Cross References report."
    )
)


# ============================================================
# INPUT MODE
# ============================================================

input_mode = st.radio(
    "Input method",
    [
        "Manual Search",
        "Upload Excel"
    ],
    horizontal=True,
    label_visibility="collapsed"
)


prepared_rows = []


# ============================================================
# MANUAL SEARCH
# ============================================================

if input_mode == "Manual Search":

    left, right = st.columns(
        [3.4, 1]
    )


    with left:

        customer_input = st.text_area(
            "Customer Part Numbers",
            value="L42508\n96337133",
            height=175,
            placeholder=(
                "Enter one customer position per line..."
            ),
            help=(
                "Several part numbers inside one customer "
                "position can be separated by commas, "
                "semicolons or slashes."
            )
        )


        manual_values = (
            customer_input.splitlines()
        )


        prepared_rows = prepare_customer_rows(
            manual_values
        )


    with right:

        source_filter = st.selectbox(
            "Cross Type",
            [
                "All",
                "OEM",
                "Aftermarket"
            ]
        )


        st.html(
            """
            <div class="search-tips">

                <div class="search-tips-title">
                    SEARCH TIPS
                </div>

                One customer position per line.<br>
                Multiple references can stay in the same line.<br>
                Spaces and hyphens are normalized automatically.

            </div>
            """
        )


# ============================================================
# EXCEL UPLOAD
# ============================================================

else:

    source_filter = st.selectbox(
        "Cross Type",
        [
            "All",
            "OEM",
            "Aftermarket"
        ]
    )


    uploaded_file = st.file_uploader(
        "Upload Customer Excel File",
        type=["xlsx"]
    )


    if uploaded_file is not None:

        try:

            excel_file = pd.ExcelFile(
                uploaded_file,
                engine="openpyxl"
            )


            col1, col2 = st.columns(2)


            with col1:

                selected_sheet = st.selectbox(
                    "1. Select Sheet",
                    excel_file.sheet_names
                )


            customer_excel = pd.read_excel(
                excel_file,
                sheet_name=selected_sheet,
                dtype=str
            )


            customer_excel.columns = [
                str(column).strip()
                for column in customer_excel.columns
            ]


            if customer_excel.empty:

                st.warning(
                    "The selected sheet is empty."
                )


            else:

                with col2:

                    selected_column = st.selectbox(
                        "2. Customer Part Number Column",
                        customer_excel.columns.tolist()
                    )


                selected_values = (
                    customer_excel[
                        selected_column
                    ].tolist()
                )


                prepared_rows = prepare_customer_rows(
                    selected_values
                )


                non_empty_count = sum(
                    1
                    for item in prepared_rows
                    if item["raw_value"]
                )


                st.html(
                    f"""
                    <div class="upload-ready">
                        <strong>{non_empty_count:,}</strong>
                        customer positions ready for search
                    </div>
                    """
                )


                st.dataframe(
                    customer_excel[
                        [selected_column]
                    ].head(8),
                    width="stretch",
                    hide_index=True
                )


        except Exception as error:

            st.error(
                f"Unable to read Excel file: {error}"
            )


# ============================================================
# SEARCH BUTTON
# ============================================================

st.write("")


search_button = st.button(
    "SEARCH CROSS REFERENCES",
    type="primary",
    width="stretch"
)


# ============================================================
# SESSION STATE
# ============================================================

if "search_results" not in st.session_state:

    st.session_state.search_results = None


if "report_reference" not in st.session_state:

    st.session_state.report_reference = None


if "report_datetime" not in st.session_state:

    st.session_state.report_datetime = None


if "report_sales_manager" not in st.session_state:

    st.session_state.report_sales_manager = ""


# ============================================================
# RUN SEARCH
# ============================================================

if search_button:

    if not sales_manager.strip():

        st.warning(
            "Please enter Sales Manager name."
        )


    elif not prepared_rows:

        st.warning(
            "No customer data is available for search."
        )


    else:

        report_datetime = datetime.now(
            DUBAI_TIMEZONE
        )


        reference_no = (
            "CR-"
            + report_datetime.strftime(
                "%Y-%m%d-%H%M"
            )
        )


        st.session_state.report_datetime = (
            report_datetime
        )


        st.session_state.report_reference = (
            reference_no
        )


        st.session_state.report_sales_manager = (
            sales_manager.strip()
        )


        with st.spinner(
            "Searching LYNXauto database..."
        ):

            st.session_state.search_results = (
                search_crosses(
                    prepared_rows,
                    source_filter
                )
            )


# ============================================================
# RESULTS
# ============================================================

if st.session_state.search_results is not None:

    df = st.session_state.search_results


    # ========================================================
    # ADD MANUAL CROSS
    # ========================================================

    st.html(
        """
        <div class="section-title">
            <div class="section-red-line"></div>
            Add Manual Cross
    </div>
    """
)


    customer_positions = (
            df[
                (
                    df["Customer Part Number"]
                    .fillna("")
                    .astype(str)
                    .str.strip()
                    != ""
                )
                &
                (
                    df["Status"]
                    .fillna("")
                    .astype(str)
                    .str.strip()
                    == "Not Found"
                )
            ][
                [
                    "No.",
                    "Customer Part Number"
                ]
            ]
            .drop_duplicates(
                subset=["No."]
            )
            .sort_values("No.")
        )


    position_options = {}


    for _, position_row in (
        customer_positions.iterrows()
    ):

        row_no = int(
            position_row["No."]
        )

        customer_number = str(
            position_row[
                "Customer Part Number"
            ]
        )


        label = (
            f"{row_no}. "
            f"{customer_number}"
        )


        position_options[
            label
        ] = row_no


    if position_options:

        manual_col1, manual_col2, manual_col3 = (
            st.columns(
                [2.2, 1.4, 2.4]
            )
        )


        with manual_col1:

            selected_position_label = (
                st.selectbox(
                    "Customer Position",
                    list(
                        position_options.keys()
                    ),
                    key="manual_customer_position"
                )
            )


            with manual_col2:

                manual_lynx_number = (
                    st.text_input(
                        "LYNXauto Number",
                        placeholder="e.g. G32877LR",
                        key="manual_lynx_number",
                        on_change=update_manual_description
                    )
                )


                with manual_col3:

                    manual_description = (
                        st.text_input(
                            "Description (optional)",
                            placeholder=(
                                "Enter product description..."
                            ),
                            key="manual_description"
                        )
                    )


        if manual_lynx_number:

            if st.session_state.get(
                "manual_description_found",
                False
            ):

                st.success(
                    "Product found in LYNX database."
                )

            else:

                st.info(
                    "LYNX number not found in database. "
                    "Enter description manually."
                )

        add_manual_button = st.button(
            "ADD MANUAL CROSS",
            width="stretch",
            key="add_manual_cross_button"
        )


        if add_manual_button:

            selected_row_no = (
                position_options[
                    selected_position_label
                ]
            )


            (
                updated_results,
                success,
                message
            ) = add_manual_cross(
                st.session_state.search_results,
                selected_row_no,
                manual_lynx_number,
                manual_description
            )


            if success:

                st.session_state.search_results = (
                    updated_results
                )

                st.success(
                    message
                )

                st.rerun()


            else:

                st.warning(
                    message
                )


    else:

        st.success(
            "All customer positions have matches."
        )

        # ========================================================
    # REMOVE MANUAL CROSS
    # ========================================================

    df = st.session_state.search_results


    manual_rows = df[
        df["Cross Type"]
        .fillna("")
        .astype(str)
        .str.strip()
        == "Manual"
    ].copy()


    if not manual_rows.empty:

        st.html(
            """
            <div class="section-title">
                <div class="section-red-line"></div>
                Remove Manual Cross
            </div>
            """
        )


        manual_cross_options = {}


        for _, manual_row in (
            manual_rows.iterrows()
        ):

            row_no = int(
                manual_row["No."]
            )

            customer_number = str(
                manual_row[
                    "Customer Part Number"
                ]
            )

            lynx_number = str(
                manual_row[
                    "LYNXauto Number"
                ]
            )


            label = (
                f"{row_no}. "
                f"{customer_number} → "
                f"{lynx_number}"
            )


            manual_cross_options[
                label
            ] = (
                row_no,
                lynx_number
            )


        remove_col1, remove_col2 = (
            st.columns(
                [4, 1]
            )
        )


        with remove_col1:

            selected_manual_cross = (
                st.selectbox(
                    "Manual Cross",
                    list(
                        manual_cross_options.keys()
                    ),
                    key="remove_manual_cross_select"
                )
            )


        with remove_col2:

            st.write("")

            st.write("")

            remove_manual_button = st.button(
                "REMOVE",
                width="stretch",
                key="remove_manual_cross_button"
            )


        if remove_manual_button:

            (
                remove_row_no,
                remove_lynx_number
            ) = manual_cross_options[
                selected_manual_cross
            ]


            (
                updated_results,
                success,
                message
            ) = remove_manual_cross(
                st.session_state.search_results,
                remove_row_no,
                remove_lynx_number
            )


            if success:

                st.session_state.search_results = (
                    updated_results
                )

                st.success(
                    message
                )

                st.rerun()


            else:

                st.warning(
                    message
                )


    # Refresh dataframe
    df = st.session_state.search_results


    # ========================================================
    # COUNTERS
    # ========================================================

    found_count = len(
        df[
            df["Status"] == "Found"
        ]
    )


    not_found_count = len(
        df[
            df["Status"] == "Not Found"
        ]
    )


    customer_rows_count = (
        df[
            df["Status"] != ""
        ]["No."]
        .nunique()
    )


    # ========================================================
    # RESULTS TITLE
    # ========================================================

    st.html(
        """
        <div class="section-title">
            <div class="section-red-line"></div>
            Search Results
        </div>
        """
    )


    # ========================================================
    # METRICS
    # ========================================================

    m1, m2, m3 = st.columns(3)


    m1.metric(
        "Customer Positions",
        customer_rows_count
    )


    m2.metric(
        "Crosses Found",
        found_count
    )


    m3.metric(
        "Not Found",
        not_found_count
    )


    # ========================================================
    # TOOLBAR
    # ========================================================

    filter_col, search_col, download_col = st.columns(
        [1.1, 2.4, 1]
    )


    with filter_col:

        result_status_filter = st.selectbox(
            "Result Status",
            [
                "All",
                "Found",
                "Not Found"
            ],
            key="result_status_filter"
        )


    with search_col:

        result_search = st.text_input(
            "Search Results",
            placeholder=(
                "Customer PN, LYNX number, description..."
            ),
            key="result_search"
        )


    # ========================================================
    # FILTER RESULTS
    # ========================================================

    filtered_df = filter_results(
        df,
        result_status_filter,
        result_search
    )


    # ========================================================
    # EXCEL DOWNLOAD
    # ========================================================

    excel_data = create_excel_download(
        df,
        st.session_state.report_sales_manager,
        st.session_state.report_reference,
        st.session_state.report_datetime
    )


    safe_sales_manager = re.sub(
        r'[\\/:*?"<>|]',
        "",
        st.session_state.report_sales_manager
    ).strip()


    if not safe_sales_manager:

        safe_sales_manager = "Sales Manager"


    excel_file_name = (
        "Cross References - "
        f"{safe_sales_manager} - "
        f"{st.session_state.report_reference}.xlsx"
    )


    with download_col:

        st.write("")


        st.download_button(
            label="DOWNLOAD XLSX",
            data=excel_data,
            file_name=excel_file_name,
            mime=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            width="stretch"
        )


    # ========================================================
    # RESULT COUNT
    # ========================================================

    st.html(
        f"""
        <div class="results-count">
            Showing <strong>{len(filtered_df):,}</strong>
            of <strong>{len(df):,}</strong> result rows
        </div>
        """
    )


    # ========================================================
    # TABLE HEIGHT
    # ========================================================

    table_height = min(
        max(
            180,
            42 + len(filtered_df) * 35
        ),
        520
    )


    # ========================================================
    # BUILD CLICKABLE TABLE
    # ========================================================

    display_table = build_display_table(
        filtered_df
    )


    # ========================================================
    # TABLE
    # ========================================================

    st.dataframe(
        display_table,
        width="stretch",
        height=table_height,
        hide_index=True,

        column_order=[
            "No.",
            "Customer Part Number",
            "LYNXauto Number",
            "Description",
            "Cross Type",
            "Status"
        ],

        column_config={

            "No.":
                st.column_config.NumberColumn(
                    "No.",
                    width="small"
                ),

            "Customer Part Number":
                st.column_config.TextColumn(
                    "Customer Part Number",
                    width="large"
                ),

            "LYNXauto Number":
                st.column_config.LinkColumn(
                    "LYNXauto Number",
                    width="medium",
                    help=(
                        "Click the LYNX number to open "
                        "the official LYNX e-catalogue"
                    ),
                    display_text=(
                        r"https://ecatalogue\.lynxauto\.jp/(.*?)\.html"
                    )
                ),

            "Description":
                st.column_config.TextColumn(
                    "Description",
                    width="large"
                ),

            "Cross Type":
                st.column_config.TextColumn(
                    "Cross Type",
                    width="small"
                ),

            "Status":
                st.column_config.TextColumn(
                    "Status",
                    width="small"
                )
        }
    )


# ============================================================
# FOOTER
# ============================================================

st.html(
    """
    <div class="lynx-footer">
        LYNX Cross Search • Internal Product Tool
    </div>
    """
)